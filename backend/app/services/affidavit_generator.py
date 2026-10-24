import io
import re
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from app.models.affidavit import AffidavitRequest, parse_affidavit_from_query


def _apply_standard_page_setup(doc: Document):
    """
    Configures 1-inch margins and strictly enforces Times New Roman font,
    12pt body size, and justified paragraph alignment.
    """
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal_style.paragraph_format.line_spacing = 1.2

    # Set explicit font elements in Word XML for ASCII, High-ANSI, and Complex Scripts
    rPr = normal_style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")


def _append_signatures_table(
    doc: Document,
    deponent_name: str,
    former_title: str = "Formerly known as",
    former_value: str = "",
):
    """Renders a standard 2-column legal attestation and deponent signature block."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(3.25), Inches(3.25)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)

    # Left Cell: Attestation / Notary
    lp1 = left_cell.paragraphs[0]
    lp1.paragraph_format.space_after = Pt(2)
    r1 = lp1.add_run("Identified by me:\n\n\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.italic = True
    r2 = lp1.add_run("_______________________________\n")
    r2.font.name = "Times New Roman"
    r3 = lp1.add_run("Advocate\n\n\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(11)

    lp2 = left_cell.add_paragraph()
    lp2.paragraph_format.space_after = Pt(2)
    r4 = lp2.add_run("BEFORE ME\n\n\n")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(11)
    r4.bold = True
    r5 = lp2.add_run("_______________________________\n")
    r5.font.name = "Times New Roman"
    r6 = lp2.add_run("NOTARY PUBLIC / OATH COMMISSIONER\n")
    r6.font.name = "Times New Roman"
    r6.font.size = Pt(10)
    r6.bold = True
    r7 = lp2.add_run("Office Seal & Registration No.:")
    r7.font.name = "Times New Roman"
    r7.font.size = Pt(10)

    # Right Cell: Deponent
    rp1 = right_cell.paragraphs[0]
    rp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp1.paragraph_format.space_after = Pt(4)
    r8 = rp1.add_run("\n\n\n\n_______________________________\n")
    r8.font.name = "Times New Roman"
    r_dep = rp1.add_run("DEPONENT\n")
    r_dep.font.name = "Times New Roman"
    r_dep.font.size = Pt(12)
    r_dep.bold = True
    r9 = rp1.add_run(f"({deponent_name.strip()})\n")
    r9.font.name = "Times New Roman"
    r9.font.size = Pt(12)
    r9.bold = True
    if former_value and former_value.strip():
        r10 = rp1.add_run(f"{former_title}: {former_value.strip()}\n")
        r10.font.name = "Times New Roman"
        r10.font.size = Pt(10)
    r11 = rp1.add_run("Signature: ____________________\n")
    r11.font.name = "Times New Roman"
    r11.font.size = Pt(11)


def _append_verification(doc: Document, city: str, state: str, clause_count: int = 6):
    """Appends standard sworn verification block strictly adhering to 14pt heading and justified alignment."""
    verif_heading = doc.add_paragraph()
    verif_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    verif_heading.paragraph_format.space_before = Pt(14)
    verif_heading.paragraph_format.space_after = Pt(6)
    vh_run = verif_heading.add_run("VERIFICATION")
    vh_run.font.name = "Times New Roman"
    vh_run.font.size = Pt(14)
    vh_run.bold = True
    vh_run.underline = True

    verif_p = doc.add_paragraph()
    verif_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    verif_p.paragraph_format.line_spacing = 1.2
    verif_p.paragraph_format.space_after = Pt(12)
    vp_run = verif_p.add_run(
        f"I, the above-named deponent, do hereby verify and declare on oath that the contents of paragraphs 1 to {clause_count} "
        "above are true and correct to the best of my personal knowledge, belief, and records. No part of it is false, "
        "and nothing material has been concealed or misrepresented therefrom."
    )
    vp_run.font.name = "Times New Roman"
    vp_run.font.size = Pt(12)

    date_place_p = doc.add_paragraph()
    date_place_p.paragraph_format.space_after = Pt(20)
    dp_run = date_place_p.add_run(
        f"Verified at {city.strip()}, {state.strip()} on this _____ day of __________________, 20____."
    )
    dp_run.font.name = "Times New Roman"
    dp_run.font.size = Pt(12)
    dp_run.italic = True


# -------------------------------------------------------------------------
# 1. Surname Change / Last Name Change Affidavit
# -------------------------------------------------------------------------
def generate_surname_change_affidavit_docx(data: AffidavitRequest) -> io.BytesIO:
    doc = Document()
    _apply_standard_page_setup(doc)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("AFFIDAVIT FOR CHANGE OF SURNAME / LAST NAME")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.bold = True
    run_title.underline = True

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run(
        "(To be executed on non-judicial stamp paper of prescribed value and attested before a Notary Public / Oath Commissioner)"
    )
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.italic = True

    intro_p = doc.add_paragraph()
    intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_p.paragraph_format.line_spacing = 1.25
    intro_p.paragraph_format.space_after = Pt(10)

    r_intro1 = intro_p.add_run("I, ")
    r_intro1.font.name = "Times New Roman"
    r_name = intro_p.add_run(str(data.full_name).strip())
    r_name.font.name = "Times New Roman"
    r_name.bold = True

    r_rel = intro_p.add_run(", Son / Daughter / Spouse of ")
    r_rel.font.name = "Times New Roman"
    r_parent = intro_p.add_run(str(data.father_or_spouse_name).strip())
    r_parent.font.name = "Times New Roman"
    r_parent.bold = True

    r_age = intro_p.add_run(f", aged about {str(data.age).strip()} years, residing at ")
    r_age.font.name = "Times New Roman"
    r_addr = intro_p.add_run(f"{str(data.address).strip()}, {str(data.city).strip()}, {str(data.state).strip()}")
    r_addr.font.name = "Times New Roman"
    r_addr.bold = True

    r_end = intro_p.add_run(", do hereby solemnly affirm and state on oath as follows:")
    r_end.font.name = "Times New Roman"

    old_display = (data.old_value or data.full_name).strip()
    new_display = (data.new_value or data.new_name).strip()

    # Determine old and new surname tokens
    old_surname = old_display.split()[-1] if " " in old_display else old_display
    new_surname = new_display.split()[-1] if " " in new_display else new_display

    clauses = [
        [
            ("1. ", False),
            ("That my full name with my original surname is currently recorded across my educational documents, identification proofs, and official government records as ", False),
            (f'"{old_display}"', True),
            (".", False),
        ],
        [
            ("2. ", False),
            ("That out of my free will, bona fide personal choice, and legal decision, I have formally resolved to abandon, renounce, and relinquish my former surname ", False),
            (f'"{old_surname}"', True),
            (" and to officially adopt and assume the new surname ", False),
            (f'"{new_surname}"', True),
            (", such that my complete lawful name henceforth shall be ", False),
            (f'"{new_display}"', True),
            (".", False),
        ],
        [
            ("3. ", False),
            ("That henceforth, for all official, statutory, legal, administrative, educational, financial, employment, passport, and personal purposes, I shall be designated, known, and addressed solely as ", False),
            (f'"{new_display}"', True),
            (".", False),
        ],
        [
            ("4. ", False),
            ("That both names, namely my former name ", False),
            (f'"{old_display}"', True),
            (" and my newly adopted name with surname ", False),
            (f'"{new_display}"', True),
            (", designate and refer to one and the same living person, that is, myself, the deponent herein.", False),
        ],
        [
            ("5. ", False),
            ("That I shall at all times hereafter in all deeds, records, writings, transactions, and legal proceedings sign and use the name ", False),
            (f'"{new_display}"', True),
            (" in full replacement and substitution of my former name ", False),
            (f'"{old_display}"', True),
            (".", False),
        ],
        [
            ("6. ", False),
            ("That this affidavit is executed for placing the change of surname on official record and for submitting the same before UIDAI (Aadhaar), Passport Seva Kendra, Income Tax Department (PAN), Banking Institutions, Gazette notification authorities, and Educational Boards/Universities for necessary records updating.", False),
        ],
    ]

    for clause in clauses:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.line_spacing = 1.2
        cp.paragraph_format.space_after = Pt(8)
        for text, is_bold in clause:
            run = cp.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            if is_bold:
                run.bold = True

    _append_verification(doc, data.city, data.state, clause_count=6)
    _append_signatures_table(doc, new_display, "Formerly known as", old_display)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------------
# 2. Name Change / Spelling Correction
# -------------------------------------------------------------------------
def generate_name_change_affidavit_docx(data: AffidavitRequest) -> io.BytesIO:
    doc = Document()
    _apply_standard_page_setup(doc)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("AFFIDAVIT FOR CHANGE OF NAME / SPELLING CORRECTION")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.bold = True
    run_title.underline = True

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run(
        "(To be executed on non-judicial stamp paper of prescribed value and attested before a Notary Public / Oath Commissioner)"
    )
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.italic = True

    intro_p = doc.add_paragraph()
    intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_p.paragraph_format.line_spacing = 1.25
    intro_p.paragraph_format.space_after = Pt(10)

    intro_p.add_run("I, ")
    r_name = intro_p.add_run(str(data.full_name).strip())
    r_name.bold = True
    intro_p.add_run(", Son / Daughter / Spouse of ")
    r_parent = intro_p.add_run(str(data.father_or_spouse_name).strip())
    r_parent.bold = True
    intro_p.add_run(f", aged about {str(data.age).strip()} years, residing at ")
    r_addr = intro_p.add_run(f"{str(data.address).strip()}, {str(data.city).strip()}, {str(data.state).strip()}")
    r_addr.bold = True
    intro_p.add_run(", do hereby solemnly affirm and state on oath as follows:")

    for r in intro_p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    old_display = (data.old_value or data.full_name).strip()
    new_display = (data.new_value or data.new_name).strip()

    clauses = [
        [
            ("1. ", False),
            ("That my name is currently recorded in my identification proofs, educational documents, and official records as ", False),
            (f'"{old_display}"', True),
            (".", False),
        ],
        [
            ("2. ", False),
            ("That due to typographical/spelling discrepancy and personal choice, I have decided to officially change and correct the spelling/record of my name from ", False),
            (f'"{old_display}"', True),
            (" to ", False),
            (f'"{new_display}"', True),
            (".", False),
        ],
        [
            ("3. ", False),
            ("That henceforth, for all official, legal, administrative, educational, financial, and personal purposes, I shall be known, designated, and addressed solely as ", False),
            (f'"{new_display}"', True),
            (".", False),
        ],
        [
            ("4. ", False),
            ("That both the names, namely ", False),
            (f'"{old_display}"', True),
            (" and ", False),
            (f'"{new_display}"', True),
            (", refer to one and the same person, that is, myself, the deponent.", False),
        ],
        [
            ("5. ", False),
            ("That I shall at all times hereafter in all deeds, writings, transactions, records, and legal proceedings sign and use the name ", False),
            (f'"{new_display}"', True),
            (" in substitution and replacement of my earlier name ", False),
            (f'"{old_display}"', True),
            (".", False),
        ],
        [
            ("6. ", False),
            ("That this affidavit is executed for the purpose of placing the change/correction of my name on record and for submitting the same before all concerned government authorities, Gazette notification offices, UIDAI (Aadhaar), Passport Seva, Income Tax Department (PAN), Banking Institutions, and Educational Boards/Universities for necessary records updating.", False),
        ],
    ]

    for clause in clauses:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.line_spacing = 1.2
        cp.paragraph_format.space_after = Pt(8)
        for text, is_bold in clause:
            run = cp.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            if is_bold:
                run.bold = True

    _append_verification(doc, data.city, data.state, clause_count=6)
    _append_signatures_table(doc, new_display, "Formerly known as", old_display)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------------
# 3. Address Change / Proof of Residence
# -------------------------------------------------------------------------
def generate_address_change_affidavit_docx(data: AffidavitRequest) -> io.BytesIO:
    doc = Document()
    _apply_standard_page_setup(doc)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("AFFIDAVIT FOR CHANGE OF RESIDENTIAL ADDRESS / PROOF OF RESIDENCE")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.bold = True
    run_title.underline = True

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run(
        "(To be executed on non-judicial stamp paper of prescribed value and attested before a Notary Public / Oath Commissioner)"
    )
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.italic = True

    intro_p = doc.add_paragraph()
    intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_p.paragraph_format.line_spacing = 1.25
    intro_p.paragraph_format.space_after = Pt(10)

    intro_p.add_run("I, ")
    r_name = intro_p.add_run(str(data.full_name).strip())
    r_name.bold = True
    intro_p.add_run(", Son / Daughter / Spouse of ")
    r_parent = intro_p.add_run(str(data.father_or_spouse_name).strip())
    r_parent.bold = True
    intro_p.add_run(f", aged about {str(data.age).strip()} years, currently residing at ")
    r_addr = intro_p.add_run(f"{str(data.new_value or data.address).strip()}, {str(data.city).strip()}, {str(data.state).strip()}")
    r_addr.bold = True
    intro_p.add_run(", do hereby solemnly affirm and state on oath as follows:")

    for r in intro_p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    old_addr = (data.old_value or "[Previous Address]").strip()
    new_addr = (data.new_value or data.address or "[New Address]").strip()

    clauses = [
        [
            ("1. ", False),
            ("That I am a bona fide citizen and permanent resident of India.", False),
        ],
        [
            ("2. ", False),
            ("That my residential address was previously recorded in my official documents, records, and identification proofs as ", False),
            (f'"{old_addr}"', True),
            (".", False),
        ],
        [
            ("3. ", False),
            ("That I have recently shifted and relocated my residence from the aforesaid previous address to my current residence at ", False),
            (f'"{new_addr}, {data.city.strip()}, {data.state.strip()}"', True),
            (", where I am residing peacefully with my family members.", False),
        ],
        [
            ("4. ", False),
            ("That the said new address, namely ", False),
            (f'"{new_addr}, {data.city.strip()}, {data.state.strip()}"', True),
            (", is my true, correct, and actual present residential address for all administrative, postal, financial, legal, and government correspondence.", False),
        ],
        [
            ("5. ", False),
            ("That I undertake to notify all relevant statutory authorities, educational institutions, and financial bodies regarding this change of residential address.", False),
        ],
        [
            ("6. ", False),
            ("That this affidavit is executed for the purpose of placing on official record my changed residential address and for submitting the same before UIDAI (Aadhaar), Passport Seva, Election Commission (Voter ID), Regional Transport Office (Driving License), and Banking Institutions for updating their records accordingly.", False),
        ],
    ]

    for clause in clauses:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.line_spacing = 1.2
        cp.paragraph_format.space_after = Pt(8)
        for text, is_bold in clause:
            run = cp.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            if is_bold:
                run.bold = True

    _append_verification(doc, data.city, data.state, clause_count=6)
    _append_signatures_table(doc, data.full_name, "Previous Address", old_addr)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------------
# 4. Date of Birth Correction
# -------------------------------------------------------------------------
def generate_dob_correction_affidavit_docx(data: AffidavitRequest) -> io.BytesIO:
    doc = Document()
    _apply_standard_page_setup(doc)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("AFFIDAVIT FOR CORRECTION OF DATE OF BIRTH")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.bold = True
    run_title.underline = True

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run(
        "(To be executed on non-judicial stamp paper of prescribed value and attested before a Notary Public / Oath Commissioner)"
    )
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.italic = True

    intro_p = doc.add_paragraph()
    intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_p.paragraph_format.line_spacing = 1.25
    intro_p.paragraph_format.space_after = Pt(10)

    intro_p.add_run("I, ")
    r_name = intro_p.add_run(str(data.full_name).strip())
    r_name.bold = True
    intro_p.add_run(", Son / Daughter / Spouse of ")
    r_parent = intro_p.add_run(str(data.father_or_spouse_name).strip())
    r_parent.bold = True
    intro_p.add_run(f", aged about {str(data.age).strip()} years, residing at ")
    r_addr = intro_p.add_run(f"{str(data.address).strip()}, {str(data.city).strip()}, {str(data.state).strip()}")
    r_addr.bold = True
    intro_p.add_run(", do hereby solemnly affirm and state on oath as follows:")

    for r in intro_p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    old_dob = (data.old_value or "[Incorrect Recorded DOB]").strip()
    new_dob = (data.new_value or "[Correct Actual DOB]").strip()

    clauses = [
        [
            ("1. ", False),
            ("That I am a bona fide citizen and permanent resident of India.", False),
        ],
        [
            ("2. ", False),
            ("That in my official documents, educational marksheets, and identity records, my date of birth has been erroneously and mistakenly recorded as ", False),
            (f'"{old_dob}"', True),
            (".", False),
        ],
        [
            ("3. ", False),
            ("That my genuine, actual, and lawful date of birth is ", False),
            (f'"{new_dob}"', True),
            (", which is duly substantiated and confirmed by my authentic civil birth certificate and hospital records.", False),
        ],
        [
            ("4. ", False),
            ("That the earlier recorded date of birth, namely ", False),
            (f'"{old_dob}"', True),
            (", occurred purely due to an inadvertent typographical and clerical oversight at the time of entry.", False),
        ],
        [
            ("5. ", False),
            ("That henceforth, my true and authentic date of birth for all legal, educational, employment, passport, banking, and government identification records shall be strictly acknowledged as ", False),
            (f'"{new_dob}"', True),
            (".", False),
        ],
        [
            ("6. ", False),
            ("That this affidavit is executed for the purpose of requesting all concerned statutory authorities, Educational Boards/Universities, UIDAI (Aadhaar), Passport Seva Kendra, and Employers to correct and update my date of birth from ", False),
            (f'"{old_dob}"', True),
            (" to ", False),
            (f'"{new_dob}"', True),
            (" in all official registers and records.", False),
        ],
    ]

    for clause in clauses:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.line_spacing = 1.2
        cp.paragraph_format.space_after = Pt(8)
        for text, is_bold in clause:
            run = cp.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            if is_bold:
                run.bold = True

    _append_verification(doc, data.city, data.state, clause_count=6)
    _append_signatures_table(doc, data.full_name, "Incorrect Recorded DOB", old_dob)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------------
# 5. Relation Update (Parent / Spouse Name Correction)
# -------------------------------------------------------------------------
def generate_relation_update_affidavit_docx(data: AffidavitRequest) -> io.BytesIO:
    doc = Document()
    _apply_standard_page_setup(doc)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("AFFIDAVIT FOR CORRECTION / UPDATE OF PARENT OR SPOUSE NAME")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.bold = True
    run_title.underline = True

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run(
        "(To be executed on non-judicial stamp paper of prescribed value and attested before a Notary Public / Oath Commissioner)"
    )
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.italic = True

    intro_p = doc.add_paragraph()
    intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_p.paragraph_format.line_spacing = 1.25
    intro_p.paragraph_format.space_after = Pt(10)

    intro_p.add_run("I, ")
    r_name = intro_p.add_run(str(data.full_name).strip())
    r_name.bold = True
    intro_p.add_run(", Son / Daughter / Spouse of ")
    r_parent = intro_p.add_run(str(data.new_value or data.father_or_spouse_name).strip())
    r_parent.bold = True
    intro_p.add_run(f", aged about {str(data.age).strip()} years, residing at ")
    r_addr = intro_p.add_run(f"{str(data.address).strip()}, {str(data.city).strip()}, {str(data.state).strip()}")
    r_addr.bold = True
    intro_p.add_run(", do hereby solemnly affirm and state on oath as follows:")

    for r in intro_p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    old_rel = (data.old_value or "[Incorrect Parent/Spouse Name]").strip()
    new_rel = (data.new_value or data.father_or_spouse_name or "[Correct Parent/Spouse Name]").strip()

    clauses = [
        [
            ("1. ", False),
            ("That I am a bona fide citizen and permanent resident of India residing at the address stated above.", False),
        ],
        [
            ("2. ", False),
            ("That in some of my educational records, identification proofs, and official documents, my father's / mother's / spouse's name has been recorded as ", False),
            (f'"{old_rel}"', True),
            (".", False),
        ],
        [
            ("3. ", False),
            ("That due to spelling variations and clerical typographical oversight, the said name was recorded as ", False),
            (f'"{old_rel}"', True),
            (", whereas the true, legal, and authentic name is ", False),
            (f'"{new_rel}"', True),
            (".", False),
        ],
        [
            ("4. ", False),
            ("That both names, namely ", False),
            (f'"{old_rel}"', True),
            (" and ", False),
            (f'"{new_rel}"', True),
            (", pertain to one and the same person, who is my biological parent / lawful spouse, and there is no other person by that name.", False),
        ],
        [
            ("5. ", False),
            ("That henceforth, my parent's / spouse's name should be recognized, addressed, and entered solely as ", False),
            (f'"{new_rel}"', True),
            (" for all official, statutory, educational, and legal records.", False),
        ],
        [
            ("6. ", False),
            ("That this affidavit is executed to request UIDAI (Aadhaar), Passport Seva, Educational Boards/Universities, and Employers to record the corrected name ", False),
            (f'"{new_rel}"', True),
            (" in substitution of the former entry ", False),
            (f'"{old_rel}"', True),
            (".", False),
        ],
    ]

    for clause in clauses:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.line_spacing = 1.2
        cp.paragraph_format.space_after = Pt(8)
        for text, is_bold in clause:
            run = cp.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            if is_bold:
                run.bold = True

    _append_verification(doc, data.city, data.state, clause_count=6)
    _append_signatures_table(doc, data.full_name, "Formerly Recorded Relative Name", old_rel)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------------
# 6. One and the Same Person Affidavit
# -------------------------------------------------------------------------
def generate_one_and_same_affidavit_docx(data: AffidavitRequest) -> io.BytesIO:
    doc = Document()
    _apply_standard_page_setup(doc)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("AFFIDAVIT FOR ONE AND THE SAME PERSON")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.bold = True
    run_title.underline = True

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run(
        "(To be executed on non-judicial stamp paper of prescribed value and attested before a Notary Public / Oath Commissioner)"
    )
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.italic = True

    intro_p = doc.add_paragraph()
    intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_p.paragraph_format.line_spacing = 1.25
    intro_p.paragraph_format.space_after = Pt(10)

    intro_p.add_run("I, ")
    r_name = intro_p.add_run(str(data.full_name).strip())
    r_name.bold = True
    intro_p.add_run(", Son / Daughter / Spouse of ")
    r_parent = intro_p.add_run(str(data.father_or_spouse_name).strip())
    r_parent.bold = True
    intro_p.add_run(f", aged about {str(data.age).strip()} years, residing at ")
    r_addr = intro_p.add_run(f"{str(data.address).strip()}, {str(data.city).strip()}, {str(data.state).strip()}")
    r_addr.bold = True
    intro_p.add_run(", do hereby solemnly affirm and state on oath as follows:")

    for r in intro_p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    alias1 = (data.old_value or data.full_name).strip()
    alias2 = (data.new_value or data.new_name).strip()

    clauses = [
        [
            ("1. ", False),
            ("That I am a bona fide citizen and permanent resident of India residing at the aforesaid address.", False),
        ],
        [
            ("2. ", False),
            ("That across my various official documents, educational certificates, banking records, and identity cards, my name has been documented with slight variations, specifically as ", False),
            (f'"{alias1}"', True),
            (" and ", False),
            (f'"{alias2}"', True),
            (".", False),
        ],
        [
            ("3. ", False),
            ("That I solemnly declare on oath that both names, namely ", False),
            (f'"{alias1}"', True),
            (" and ", False),
            (f'"{alias2}"', True),
            (", designate, belong, and pertain to one and the same living person, that is, myself, the deponent herein.", False),
        ],
        [
            ("4. ", False),
            ("That there is no other or second individual by either of the aforesaid names in my family or locality, and both names signify my singular legal personality.", False),
        ],
        [
            ("5. ", False),
            ("That I shall sign, be addressed by, and use both names as required, without prejudice to my legal rights, liabilities, and identity.", False),
        ],
        [
            ("6. ", False),
            ("That this affidavit is executed to satisfy the identity verification requirements of Passport Seva, UIDAI (Aadhaar), Visa & Immigration authorities, Banks, and Universities, confirming that ", False),
            (f'"{alias1}"', True),
            (" and ", False),
            (f'"{alias2}"', True),
            (" are one and the same person.", False),
        ],
    ]

    for clause in clauses:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.line_spacing = 1.2
        cp.paragraph_format.space_after = Pt(8)
        for text, is_bold in clause:
            run = cp.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            if is_bold:
                run.bold = True

    _append_verification(doc, data.city, data.state, clause_count=6)
    _append_signatures_table(doc, data.full_name, "Alias / Name Variation", f"{alias1} / {alias2}")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------------
# Main Dispatcher & Helper Functions
# -------------------------------------------------------------------------
def generate_affidavit_docx(data: AffidavitRequest) -> io.BytesIO:
    """
    Main dispatcher: dynamically generates the appropriate Word document (.docx)
    based on the requested `document_type`.
    """
    doc_type = (data.document_type or "name_change").strip().lower()
    if "surname" in doc_type or "last_name" in doc_type or "family_name" in doc_type:
        return generate_surname_change_affidavit_docx(data)
    elif "dob" in doc_type or "birth" in doc_type:
        return generate_dob_correction_affidavit_docx(data)
    elif "relation" in doc_type or "parent" in doc_type or "spouse" in doc_type:
        return generate_relation_update_affidavit_docx(data)
    elif "same" in doc_type or "one_and" in doc_type:
        return generate_one_and_same_affidavit_docx(data)
    elif "address" in doc_type or "residence" in doc_type:
        return generate_address_change_affidavit_docx(data)
    else:
        return generate_name_change_affidavit_docx(data)


def detect_affidavit_intent(query_or_text: str) -> str:
    """Detects exact legal affidavit intent from natural language query or message text."""
    parsed = parse_affidavit_from_query(query_or_text)
    return parsed.get("document_type", "name_change")


def extract_affidavit_parameters(query_or_text: str, default_name: Optional[str] = None) -> AffidavitRequest:
    """Extracts parameters from a query and builds a validated AffidavitRequest."""
    parsed = parse_affidavit_from_query(query_or_text, default_name=default_name)
    return AffidavitRequest(**parsed)


def detect_affidavit_intent(query_or_text: str) -> str:
    """Detects exact legal affidavit intent from natural language query or message text."""
    parsed = parse_affidavit_from_query(query_or_text)
    return parsed.get("document_type", "name_change")


def extract_affidavit_parameters(query_or_text: str, default_name: Optional[str] = None) -> AffidavitRequest:
    """Extracts parameters from a query and builds a validated AffidavitRequest."""
    parsed = parse_affidavit_from_query(query_or_text, default_name=default_name)
    return AffidavitRequest(**parsed)

